"""python3 -m unittest -v"""
import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

import daily_report as dr
from make_sample import make


def write(folder: Path, name: str, text: str, enc: str = "utf-8-sig") -> None:
    (folder / name).write_bytes(text.encode(enc))


class ParsingTest(unittest.TestCase):
    def test_shift_crosses_midnight(self):
        s = dr.Settings()
        self.assertEqual(dr.shift_of(datetime(2026, 10, 29, 8, 0), s), (date(2026, 10, 29), "早班"))
        self.assertEqual(dr.shift_of(datetime(2026, 10, 29, 19, 59), s), (date(2026, 10, 29), "早班"))
        self.assertEqual(dr.shift_of(datetime(2026, 10, 29, 20, 0), s), (date(2026, 10, 29), "夜班"))
        self.assertEqual(dr.shift_of(datetime(2026, 10, 30, 7, 59), s), (date(2026, 10, 29), "夜班"))

    def test_timestamp_formats(self):
        self.assertEqual(dr.parse_ts("2026/10/29 08:00"), datetime(2026, 10, 29, 8, 0))
        self.assertEqual(dr.parse_ts("2026-10-29T08:00:30"), datetime(2026, 10, 29, 8, 0, 30))
        self.assertIsNone(dr.parse_ts("昨天早上"))

    def test_header_aliases_any_order(self):
        cols = dr.map_headers(["NG", "Good", "Input", "Timestamp"])
        self.assertEqual(cols, {"時間": 3, "投入數": 2, "良品數": 1, "不良數": 0})


class LoadTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_big5_and_utf8_files_both_load(self):
        head = "時間,投入數,良品數,不良數,停機分鐘,停機原因\n"
        write(self.dir, "A1_x.csv", head + "2026-10-29 09:00,100,99,1,0,\n", "cp950")
        write(self.dir, "A2_x.csv", head + "2026-10-29 09:00,100,98,2,10,換線\n")
        res = dr.load_folder(self.dir)
        self.assertEqual([r.machine for r in res.records], ["A1", "A2"])
        self.assertEqual(res.issues, [])

    def test_bad_rows_become_issues_not_crashes(self):
        head = "時間,投入數,良品數,不良數,停機分鐘,停機原因\n"
        body = ("2026-10-29 09:00,100,97,2,0,\n"     # 數量對不起來
                "2026-10-29 10:00,abc,0,0,0,\n"      # 無法辨識
                ",,,,,\n"                            # 空白列，略過
                "2026-10-29 11:00,100,100,0,90,\n")  # 停機過長＋未填原因
        write(self.dir, "B1_x.csv", head + body)
        res = dr.load_folder(self.dir)
        kinds = sorted(i.kind for i in res.issues)
        self.assertEqual(kinds, ["停機未填原因", "停機過長", "數量對不起來", "資料無法辨識"])
        self.assertEqual(len(res.records), 2)

    def test_missing_required_column(self):
        write(self.dir, "C1_x.csv", "時間,投入數\n2026-10-29 09:00,100\n")
        res = dr.load_folder(self.dir)
        self.assertEqual(res.records, [])
        self.assertEqual(res.issues[0].kind, "欄位缺少")

    def test_empty_folder_raises(self):
        with self.assertRaises(FileNotFoundError):
            dr.load_folder(self.dir)


class ReportTest(unittest.TestCase):
    def test_end_to_end_with_sample(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            make(tmp / "in")
            s = dr.Settings()
            load = dr.load_folder(tmp / "in", s)
            info = dr.write_report(tmp / "out.xlsx", date(2026, 10, 29), load, s)
            self.assertEqual(info["machines"], 3)
            self.assertEqual(info["records"], 71)        # 3 台 × 24 小時，扣掉 1 筆無法辨識
            self.assertIn("M02", info["low_yield"])      # 範例裡 M02 不良率較高
            kinds = {i.kind for i in load.issues}
            self.assertTrue({"停機過長", "數量對不起來", "停機未填原因", "資料無法辨識"} <= kinds)

            wb = load_workbook(tmp / "out.xlsx")
            self.assertEqual(wb.sheetnames, ["總覽", "班別", "停機原因", "異常清單", "原始資料"])
            ws = wb["總覽"]
            self.assertEqual(ws["A4"].value, "M01")
            self.assertEqual(ws["A7"].value, "合計")
            self.assertEqual(ws["B7"].value, info["input"])
            self.assertEqual(wb["班別"].max_row, 1 + 3 * 2)  # 3 台 × 2 班

    def test_only_records_of_report_date(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            (tmp / "in").mkdir()
            write(tmp / "in", "D1_x.csv", "時間,投入數,良品數,不良數\n"
                  "2026-10-29 07:00,10,10,0\n"   # 屬於 10/28 夜班
                  "2026-10-29 09:00,20,20,0\n"   # 10/29 早班
                  "2026-10-30 07:00,30,30,0\n")  # 10/29 夜班
            s = dr.Settings()
            info = dr.write_report(tmp / "o.xlsx", date(2026, 10, 29), dr.load_folder(tmp / "in", s), s)
            self.assertEqual(info["input"], 50)


if __name__ == "__main__":
    unittest.main()
